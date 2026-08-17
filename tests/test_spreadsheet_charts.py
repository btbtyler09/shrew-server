"""Chart tabularization: a chart is a view of a table, so the extractor emits
the table — legend entries as columns, axis titles as units, chart-ness
(type/title/source) preserved in the heading. Plus xl/media picture extraction."""
import io
import os

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference, ScatterChart, Series

from app.spreadsheet_extract import extract_spreadsheet, extract_spreadsheet_media


def _wb_with(chart_builder):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    rows = [("Quarter", "Revenue", "Costs"),
            ("Q1", 10, 4), ("Q2", 20, 8), ("Q3", 30, 12)]
    for r in rows:
        ws.append(r)
    chart_builder(wb, ws)
    return wb


def _extract(wb, tmp_path):
    p = tmp_path / "t.xlsx"
    wb.save(p)
    return extract_spreadsheet(str(p), str(tmp_path))


def test_multi_series_folds_with_legend_columns_and_units(tmp_path):
    def build(wb, ws):
        ch = BarChart()
        ch.title = "Quarterly"
        ch.x_axis.title = "Quarter"
        ch.y_axis.title = "USD"
        data = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=4)
        cats = Reference(ws, min_col=1, min_row=2, max_row=4)
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        ws.add_chart(ch, "E2")
    md = _extract(_wb_with(build), tmp_path)
    assert "### Chart (bar): Quarterly — source:" in md
    # legend entries as column titles, axis title as unit suffix
    assert "| Quarter | Revenue (USD) | Costs (USD) |" in md
    assert "| Q2 | 20 | 8 |" in md
    # axis titles are IN the headers, not floating bold lines
    assert "**Y-axis:**" not in md


def test_single_series_uses_axis_titles(tmp_path):
    def build(wb, ws):
        ch = BarChart()
        ch.y_axis.title = "USD"
        ch.x_axis.title = "Quarter"
        data = Reference(ws, min_col=2, min_row=1, max_row=4)
        cats = Reference(ws, min_col=1, min_row=2, max_row=4)
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        ws.add_chart(ch, "E2")
    md = _extract(_wb_with(build), tmp_path)
    assert "| Quarter | Revenue (USD) |" in md


def test_missing_labels_fall_back_never_invent(tmp_path):
    def build(wb, ws):
        ch = BarChart()          # no title, no axis titles
        data = Reference(ws, min_col=2, min_row=2, max_row=4)  # headerless
        ch.add_data(data, titles_from_data=False)
        ws.add_chart(ch, "E2")
    md = _extract(_wb_with(build), tmp_path)
    assert "### Chart (bar): (untitled)" in md
    assert "Series 1" in md


def test_pie_tabularizes(tmp_path):
    def build(wb, ws):
        ch = PieChart()
        ch.title = "Mix"
        data = Reference(ws, min_col=2, min_row=1, max_row=4)
        cats = Reference(ws, min_col=1, min_row=2, max_row=4)
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        ws.add_chart(ch, "E2")
    md = _extract(_wb_with(build), tmp_path)
    assert "### Chart (pie): Mix" in md
    assert "| Q1 | 10 |" in md


def test_scatter_recovers_x_values_from_xval(tmp_path):
    def build(wb, ws):
        ch = ScatterChart()
        ch.x_axis.title = "Costs"
        ch.y_axis.title = "Revenue"
        xref = Reference(ws, min_col=3, min_row=2, max_row=4)
        yref = Reference(ws, min_col=2, min_row=1, max_row=4)
        ch.series.append(Series(yref, xref, title_from_data=True))
        ws.add_chart(ch, "E2")
    md = _extract(_wb_with(build), tmp_path)
    # x-values must appear as the category column, not vanish
    assert "| 4 | 10 |" in md
    assert "| Costs | Revenue" in md


def test_embedded_picture_extraction(tmp_path):
    PIL = pytest.importorskip("PIL.Image")
    from openpyxl.drawing.image import Image as XLImage
    png = tmp_path / "logo.png"
    PIL.new("RGB", (24, 24), (200, 30, 30)).save(png)
    wb = Workbook()
    ws = wb.active
    ws.title = "Front"
    ws.append(("hello",))
    ws.add_image(XLImage(str(png)), "B2")
    p = tmp_path / "img.xlsx"
    wb.save(p)
    media = extract_spreadsheet_media(str(p), str(tmp_path))
    assert len(media) == 1
    assert media[0]["sheet"] == "Front"
    assert os.path.getsize(media[0]["path"]) > 0
